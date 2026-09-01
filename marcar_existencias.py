#!/usr/bin/env python3
"""Llena DISPONIBILIDAD en el catalogo cruzando contra existencias.txt.

Regla de Alek (2026-08-29): si el modelo aparece impreso en la hoja de
inventario de Rusticos, se maneja en existencia -> EN BODEGA (5 a 7 dias).
Lo que no aparece -> BAJO PEDIDO (15 a 25 dias). No hay "sobre pedido":
lo unico que cambia es el plazo.
"""
import re
import sys
import unicodedata
import openpyxl

def norm(s):
    s = unicodedata.normalize('NFKD', str(s or ''))
    s = ''.join(c for c in s if not unicodedata.combining(c))
    s = re.sub(r'[^a-z0-9]+', ' ', s.lower()).strip()
    # el catalogo y la hoja escriben distinto lo mismo
    for a, b in [(' az ', ' azul '), (' tc ', ' terracota '), (' bco ', ' blanco '),
                 (' most ', ' mostaza '), (' fdo ', ' fondo '), (' vde ', ' verde '),
                 (' ama ', ' amarillo ')]:
        s = f' {s} '.replace(a, b).strip()
    s = re.sub(r'\s+', ' ', s)
    # "P 10" y "P10" son el mismo modelo
    return re.sub(r'\b([a-z])\s+(\d)', r'\1\2', s)

def cargar(ruta):
    grupos, act = {}, None
    for ln in open(ruta, encoding='utf-8'):
        ln = ln.split('#')[0].strip()
        if not ln:
            continue
        if ln.startswith('['):
            act = ln.strip('[]')
            grupos[act] = []
        elif act:
            grupos[act].append(norm(ln))
    return grupos

def main(xlsx, listado, salida):
    g = cargar(listado)
    wb = openpyxl.load_workbook(xlsx)
    s = wb['CATALOGO']
    h = {c.value: c.column for c in s[1]}
    C = {k: h[k] for k in ('PRODUCTO', 'CATEGORIA', 'ACABADO', 'COLOR',
                           'TIPO_PIEZA', 'PIEZA_FORMA', 'BODEGA', 'DISPONIBILIDAD',
                           'ALIAS')}

    hay = fuera = 0
    sin_cruce = set()
    for r in range(2, s.max_row + 1):
        if str(s.cell(r, C['BODEGA']).value or '').upper() != 'SI':
            continue
        cat = str(s.cell(r, C['CATEGORIA']).value or '')
        acab = str(s.cell(r, C['ACABADO']).value or '')
        color = norm(s.cell(r, C['COLOR']).value)
        prod = norm(s.cell(r, C['PRODUCTO']).value)
        pieza = norm(s.cell(r, C['TIPO_PIEZA']).value) or norm(s.cell(r, C['PIEZA_FORMA']).value)
        # la hoja usa los nombres del proveedor ("Amarillo Deslavado"), el
        # catalogo los de RAAC ("Mostaza Deslavado"). ALIAS los empata.
        # OJO: ALIAS es columna prohibida en el JSON publico — aqui solo cruza.
        alias = [norm(a) for a in re.split(r'[;,|/]', str(s.cell(r, C['ALIAS']).value or '')) if a.strip()]

        if cat.startswith('Terminacion'):
            # la hoja escribe "Bagueta Azul Deslavado"; el catalogo separa
            # COLOR="Azul" y ACABADO="Deslavado" en dos columnas.
            forma = norm(s.cell(r, C['PIEZA_FORMA']).value)
            ac = norm(acab)
            claves = set()
            for pz in (pieza, forma):
                if not pz:
                    continue
                claves.add(f'{pz} {color}')
                claves.add(f'{pz} {color} {ac}'.strip())
                claves.add(f'{pz} {ac} {color}'.strip())
            for al in alias:
                for pz in (pieza, forma):
                    if pz:
                        claves.add(f'{pz} {al}')
            en = any(norm(k) in g['TERMINACION'] for k in claves if k)
        elif 'Relieve' in cat or acab.startswith('Relieve'):
            en = any(x and (x in prod or any(x in a for a in alias)) for x in g['RELIEVE'])
        elif cat == 'Talavera Decorado':
            en = any(x and (x in prod or any(x in a for a in alias)) for x in g['DECORADO'])
        elif cat == 'Talavera':
            # el liso se identifica por acabado + color: "Azul Deslavado", "Mango Mate"
            cands = [color, f'{color} {acab}', f'{acab} {color}']
            cands += alias + [f'{a} {norm(acab)}' for a in alias]
            en = any(norm(k) in g['AZULEJO_COLOR'] for k in cands)
        else:
            en = False  # boutique y pegazulejo no tienen hoja

        s.cell(r, C['DISPONIBILIDAD']).value = 'EN BODEGA' if en else 'BAJO PEDIDO'
        if en:
            hay += 1
        else:
            fuera += 1
            sin_cruce.add((cat, acab, str(s.cell(r, C['PRODUCTO']).value)))

    wb.save(salida)
    print(f'EN BODEGA: {hay} filas · BAJO PEDIDO: {fuera} filas')
    return sin_cruce

if __name__ == '__main__':
    sc = main(sys.argv[1], sys.argv[2], sys.argv[3])
    print(f'\nModelos SIN cruce ({len(sc)}) — revisar si alguno debería estar en existencia:')
    for cat, acab, prod in sorted(sc)[:60]:
        print(f'   {cat[:26]:26} {prod}')
